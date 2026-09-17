"""Pruebas de las declaraciones de proyectos.py (sin depender de leer(), aún
NotImplementedError) y del escenario real que resuelve esta tarjeta.

TARJETA: [HU-04][BE-01] (resolver_hoja_proyectos + _comun.rellenar_celdas_combinadas)
CUBRE: HU-04 / CA-2.
"""

from __future__ import annotations

import pandas as pd
import pytest
from openpyxl import Workbook

from app.modules.ingesta.persistence.lectores import _comun
from app.modules.ingesta.persistence.lectores.proyectos import resolver_hoja_proyectos
from app.shared.errors import ArchivoInvalido
from tests.fabricas import BPIN_1, HOJA_PROYECTOS, _a_bytes, construir_proyectos


class TestResolverHojaProyectos:
    def test_encuentra_la_hoja_sin_nombre_estandar(self) -> None:
        hoja, fila_encabezado = resolver_hoja_proyectos(construir_proyectos(), "proyectos.xlsx")
        assert hoja == HOJA_PROYECTOS
        assert fila_encabezado == 0

    def test_prueba_cada_hoja_hasta_encontrar_las_columnas(self) -> None:
        # La primera hoja del libro no trae las columnas requeridas; la
        # segunda sí. resolver_hoja_proyectos no puede asumir que la
        # respuesta está en la primera hoja que abre.
        libro = Workbook()
        libro.remove(libro.active)
        libro.create_sheet("Portada")["A1"] = "esta hoja no trae columnas"
        hoja_datos = libro.create_sheet("2026")
        hoja_datos.append(["Código BPIN", "Indicador de producto"])
        hoja_datos.append([BPIN_1, "459903100"])

        hoja, fila_encabezado = resolver_hoja_proyectos(_a_bytes(libro), "proyectos.xlsx")

        assert hoja == "2026"
        assert fila_encabezado == 0

    def test_lanza_archivo_invalido_si_ninguna_hoja_tiene_las_columnas(self) -> None:
        with pytest.raises(ArchivoInvalido) as exc:
            resolver_hoja_proyectos(construir_proyectos(incluir_bpin=False), "proyectos.xlsx")
        assert exc.value.detalles["motivo"] == "hoja_no_encontrada"


def test_celdas_combinadas_de_proyecto_se_propagan_a_la_fila_de_contrato() -> None:
    """Escenario real completo: resolver la hoja, leerla, y propagar BPIN y
    nombre del proyecto desde la fila combinada hacia la fila que solo trae
    datos de contrato — la peculiaridad 6 documentada en _comun.py.
    """
    contenido = construir_proyectos()
    hoja, fila_encabezado = resolver_hoja_proyectos(contenido, "proyectos.xlsx")
    df = _comun.leer_hoja(contenido, hoja, fila_encabezado)

    # Antes de propagar: la fila de contrato llega vacía en las columnas que
    # en Excel están combinadas con la fila de proyecto de arriba.
    assert df["codigo bpin"].iloc[0] == BPIN_1
    assert pd.isna(df["codigo bpin"].iloc[1])

    propagado = _comun.rellenar_celdas_combinadas(df, ["codigo bpin", "nombre del proyecto"])

    assert propagado["codigo bpin"].tolist() == [BPIN_1, BPIN_1]
    assert propagado["nombre del proyecto"].nunique() == 1
