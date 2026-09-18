"""Pruebas de resolver_hojas y leer() en lectores/ejecucion.py.

TARJETAS: [HU-03][BE-01] (resolver_hojas), [HU-03][BE-04]/[BE-05] (leer())
CUBRE: HU-03 / CA-2 — ambas pestañas (Ejecución y Contratación) se localizan
de forma independiente: la ausencia de una no impide resolver la otra, y no
se exige subirlas por separado (es un solo archivo, dos hojas).
"""

from __future__ import annotations

import io

import pytest
from openpyxl import Workbook

from app.modules.ingesta.persistence.lectores.ejecucion import LectorEjecucion, resolver_hojas
from app.shared.errors import ArchivoInvalido
from tests.fabricas import (
    COD_B,
    HOJA_CONTRATACION,
    HOJA_EJECUCION,
    construir_ejecucion,
    construir_pdt,
)


def test_resuelve_ambas_hojas_cuando_las_dos_existen() -> None:
    hoja_ejecucion, hoja_contratacion = resolver_hojas(construir_ejecucion(), "presupuestal.xlsx")

    assert hoja_ejecucion == HOJA_EJECUCION
    assert hoja_contratacion == HOJA_CONTRATACION


def test_resuelve_ejecucion_aunque_falte_contratacion() -> None:
    contenido = construir_ejecucion(incluir_contratacion=False)

    hoja_ejecucion, hoja_contratacion = resolver_hojas(contenido, "presupuestal.xlsx")

    assert hoja_ejecucion == HOJA_EJECUCION
    assert hoja_contratacion is None


def test_resuelve_contratacion_aunque_falte_ejecucion() -> None:
    contenido = construir_ejecucion(incluir_ejecucion=False)

    hoja_ejecucion, hoja_contratacion = resolver_hojas(contenido, "presupuestal.xlsx")

    assert hoja_ejecucion is None
    assert hoja_contratacion == HOJA_CONTRATACION


def test_ninguna_pestana_no_lanza_excepcion() -> None:
    contenido = construir_ejecucion(incluir_ejecucion=False, incluir_contratacion=False)

    assert resolver_hojas(contenido, "presupuestal.xlsx") == (None, None)


class TestLeer:
    """[HU-03][BE-04]/[BE-05]: orquesta resolver_hojas (ya probado por
    separado) y cubre CA-4 (pestaña faltante, nombrada), CA-5 (archivo que
    no corresponde) y CA-6 (número y descripción del contrato para la
    matriz)."""

    def test_caso_feliz_devuelve_ejecucion_y_contratacion(self) -> None:
        resultado = LectorEjecucion().leer(
            construir_ejecucion(), "presupuestal.xlsx", vigencia=2026
        )

        assert resultado.conteos == {"ejecucion": 1, "contratacion": 1}
        assert resultado.advertencias == []
        assert resultado.filas["ejecucion"][0]["cod_indicador_producto"] == COD_B
        assert resultado.filas["contratacion"][0]["cod_indicador_producto"] == COD_B

    def test_ca6_aporta_numero_y_descripcion_del_contrato(self) -> None:
        resultado = LectorEjecucion().leer(
            construir_ejecucion(), "presupuestal.xlsx", vigencia=2026
        )

        contrato = resultado.filas["contratacion"][0]
        assert contrato["numero_contrato"] == "C-001"
        assert contrato["descripcion_contrato"] == "Mantenimiento de vías terciarias"

    def test_ca4_rechaza_y_nombra_ejecucion_cuando_falta(self) -> None:
        contenido = construir_ejecucion(incluir_ejecucion=False)

        with pytest.raises(ArchivoInvalido) as exc:
            LectorEjecucion().leer(contenido, "presupuestal.xlsx", vigencia=2026)

        assert exc.value.detalles == {"motivo": "pestana_faltante", "pestana_faltante": "EJECUCION"}

    def test_ca4_rechaza_y_nombra_contratacion_cuando_falta(self) -> None:
        contenido = construir_ejecucion(incluir_contratacion=False)

        with pytest.raises(ArchivoInvalido) as exc:
            LectorEjecucion().leer(contenido, "presupuestal.xlsx", vigencia=2026)

        assert exc.value.detalles == {
            "motivo": "pestana_faltante",
            "pestana_faltante": "CONTRATACION",
        }

    def test_ca5_rechaza_un_archivo_que_no_es_el_presupuestal(self) -> None:
        """Subir el PDT donde se espera el archivo presupuestal: ninguna de
        las dos pestañas aparece, así que es CA-5 (archivo incorrecto), no
        CA-4 (pestaña faltante)."""
        with pytest.raises(ArchivoInvalido) as exc:
            LectorEjecucion().leer(construir_pdt(), "pdt.xlsx", vigencia=2026)

        assert exc.value.detalles["motivo"] == "archivo_no_corresponde"

    def test_fila_con_codigo_de_indicador_invalido_se_descarta_con_advertencia(self) -> None:
        libro = Workbook()
        libro.remove(libro.active)
        hoja = libro.create_sheet(HOJA_EJECUCION)
        hoja.append(["CodigoIndicadorCcpet"])
        hoja.append(["12345"])  # 5 dígitos: no es normalizable a 9
        hoja_contratacion = libro.create_sheet(HOJA_CONTRATACION)
        hoja_contratacion.append(["Cod Indicador Ccpet"])
        hoja_contratacion.append([COD_B])
        buffer = io.BytesIO()
        libro.save(buffer)

        resultado = LectorEjecucion().leer(buffer.getvalue(), "raro.xlsx", vigencia=2026)

        assert resultado.filas["ejecucion"] == []
        assert resultado.conteos["ejecucion"] == 0
        assert "12345" in resultado.advertencias[0]

    def test_columnas_opcionales_ausentes_no_rechazan_el_archivo(self) -> None:
        """numero_contrato/descripcion_contrato no son obligatorias (solo CA-6)."""
        libro = Workbook()
        libro.remove(libro.active)
        libro.create_sheet(HOJA_EJECUCION).append(["CodigoIndicadorCcpet"])
        libro[HOJA_EJECUCION].append([COD_B])
        hoja_contratacion = libro.create_sheet(HOJA_CONTRATACION)
        hoja_contratacion.append(["Cod Indicador Ccpet"])
        hoja_contratacion.append([COD_B])
        buffer = io.BytesIO()
        libro.save(buffer)

        resultado = LectorEjecucion().leer(buffer.getvalue(), "raro.xlsx", vigencia=2026)

        contrato = resultado.filas["contratacion"][0]
        assert contrato["numero_contrato"] is None
        assert contrato["descripcion_contrato"] is None
