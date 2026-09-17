"""Pruebas de las declaraciones de pdt.py (sin depender de leer(), aún NotImplementedError).

TARJETAS: [HU-02][BE-01] (resolver_hoja_pdt, celdas combinadas en encabezado)
          [HU-02][BE-02] (OBLIGATORIAS)
CUBRE: HU-02 / CA-2, CA-3.

Dos huecos encontrados al comparar el texto COMPLETO de la tarjeta de Trello
de [HU-02][BE-01] contra lo entregado (no solo su docstring-resumen), y
cerrados aquí:
  1. "Si la pestaña no existe, lanza excepción identificando el problema" —
     no estaba implementado (`_comun.resolver_hoja` deliberadamente devuelve
     `None`, para permitir búsquedas independientes en HU-03). Ver
     `TestResolverHojaPdt`.
  2. "Celdas combinadas en el encabezado" — sin probar. Ver
     `test_leer_hoja_con_celda_combinada_en_el_encabezado`.
"""

from __future__ import annotations

import io

import pytest
from openpyxl import Workbook

from app.modules.ingesta.persistence.lectores import _comun
from app.modules.ingesta.persistence.lectores.pdt import (
    OBLIGATORIAS,
    alias_columna_programacion,
    resolver_hoja_pdt,
)
from app.shared.errors import ArchivoInvalido
from tests.fabricas import HOJA_PDT, construir_pdt

_ENCABEZADOS_REALES = (
    "Código de indicador de producto (MGA)",
    "Código de indicador de producto (SisPT)",
    "Producto (MGA)",
    "Indicador de Producto(MGA)",
    "Principal",
    "Programación del producto bien o servicio 2026",
    "Total 2026",
)


class TestResolverHojaPdt:
    def test_encuentra_la_pestana_cuando_existe(self) -> None:
        assert resolver_hoja_pdt(construir_pdt(), "pdt.xlsx") == HOJA_PDT

    def test_lanza_archivo_invalido_si_no_existe(self) -> None:
        libro = Workbook()
        libro.active.title = "Otra cosa"
        buffer = io.BytesIO()
        libro.save(buffer)

        with pytest.raises(ArchivoInvalido) as exc:
            resolver_hoja_pdt(buffer.getvalue(), "ejecucion.xlsx")
        assert exc.value.detalles["motivo"] == "hoja_no_encontrada"


def test_leer_hoja_con_celda_combinada_en_el_encabezado() -> None:
    """Una columna de encabezado combinada no debe tumbar ni desplazar a las
    demás: pandas la renombra a "Unnamed: N" y sigue leyendo el resto bien.
    """
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Datos"
    hoja.append(["Código de indicador de producto (MGA)", "Programación", None, "Principal"])
    hoja.merge_cells("B1:C1")
    hoja.append(["040110500", "10", None, "Sí"])
    buffer = io.BytesIO()
    libro.save(buffer)

    df = _comun.leer_hoja(buffer.getvalue(), "Datos", fila_encabezado=0)

    assert df["codigo de indicador de producto (mga)"].tolist() == ["040110500"]
    assert df["principal"].tolist() == ["Sí"]


def test_alias_de_obligatorias_coinciden_con_encabezados_reales() -> None:
    normalizados_reales = {_comun.normalizar_encabezado(c) for c in _ENCABEZADOS_REALES}
    for alias in OBLIGATORIAS.values():
        assert any(_comun.normalizar_encabezado(a) in normalizados_reales for a in alias)


def test_alias_columna_programacion_coincide_con_el_encabezado_real_de_esa_vigencia() -> None:
    generado = alias_columna_programacion(2026)
    normalizados_reales = {_comun.normalizar_encabezado(c) for c in _ENCABEZADOS_REALES}
    assert _comun.normalizar_encabezado(generado) in normalizados_reales


def test_alias_columna_programacion_cambia_con_la_vigencia() -> None:
    assert alias_columna_programacion(2025) != alias_columna_programacion(2026)
